#!/usr/bin/env python3
"""
Export sent and received emails from a Hostinger mailbox to an organised Excel workbook.

Note: SMTP can only *send* mail, it cannot read a mailbox. To export the emails
that were sent and received, this script connects to the same Hostinger account
over IMAP (imap.hostinger.com:993, SSL) using the same email address and password.

Workbook layout:
  - Summary        : totals per folder, date range, top senders / recipients
  - Received       : every message in the Inbox
  - Sent           : every message in the Sent folder
  - All Emails     : both combined, sorted newest first

Messages are fetched with BODY.PEEK so nothing is marked as read.

Usage:
  pip install openpyxl
  python export_emails.py --email you@yourdomain.com
  python export_emails.py --email you@yourdomain.com --since 2025-01-01 --output mail.xlsx

The password is read from the HOSTINGER_EMAIL_PASSWORD environment variable,
or prompted for interactively if that is not set.
"""

import argparse
import email
import getpass
import imaplib
import os
import re
import sys
from collections import Counter
from datetime import datetime
from email.header import decode_header, make_header
from email.utils import getaddresses, parseaddr, parsedate_to_datetime

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl is required: pip install openpyxl")

IMAP_HOST = "imap.hostinger.com"
IMAP_PORT = 993
BATCH_SIZE = 100
SNIPPET_LEN = 300

# Common names for the Sent folder, used if the server does not flag it with \Sent.
SENT_FALLBACKS = ["INBOX.Sent", "Sent", "Sent Items", "Sent Messages", "INBOX.Sent Items"]

COLUMNS = [
    ("Folder", 12),
    ("Date", 20),
    ("From Name", 24),
    ("From Email", 30),
    ("To", 40),
    ("Cc", 30),
    ("Subject", 50),
    ("Attachments", 8),
    ("Attachment Names", 35),
    ("Size (KB)", 10),
    ("Read", 7),
    ("Preview", 60),
    ("Message-ID", 40),
]

HEADER_FILL = PatternFill("solid", start_color="1F4E78")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=10)
TITLE_FONT = Font(name="Arial", bold=True, size=14)
BAND_FILL = PatternFill("solid", start_color="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(bottom=THIN)


# --------------------------------------------------------------------------- IMAP

def decode_str(value):
    if not value:
        return ""
    try:
        return str(make_header(decode_header(value))).strip()
    except Exception:
        return str(value).strip()


def format_addresses(value):
    pairs = getaddresses([decode_str(value)]) if value else []
    out = []
    for name, addr in pairs:
        if not addr:
            continue
        out.append(f"{name} <{addr}>" if name else addr)
    return "; ".join(out)


def parse_date(value):
    if not value:
        return None
    try:
        dt = parsedate_to_datetime(value)
    except Exception:
        return None
    if dt is None:
        return None
    if dt.tzinfo is not None:
        dt = dt.astimezone().replace(tzinfo=None)  # local time, Excel has no tz
    return dt


def extract_text_and_attachments(msg):
    text = ""
    html = ""
    attachments = []
    for part in msg.walk():
        if part.is_multipart():
            continue
        disposition = (part.get("Content-Disposition") or "").lower()
        filename = part.get_filename()
        if filename or "attachment" in disposition:
            attachments.append(decode_str(filename) or "(unnamed)")
            continue
        ctype = part.get_content_type()
        if ctype not in ("text/plain", "text/html"):
            continue
        try:
            payload = part.get_payload(decode=True) or b""
            charset = part.get_content_charset() or "utf-8"
            decoded = payload.decode(charset, errors="replace")
        except (LookupError, AttributeError):
            decoded = ""
        if ctype == "text/plain" and not text:
            text = decoded
        elif ctype == "text/html" and not html:
            html = decoded
    if not text and html:
        text = re.sub(r"(?is)<(script|style).*?</\1>", " ", html)
        text = re.sub(r"(?s)<[^>]+>", " ", text)
        text = re.sub(r"&nbsp;", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    # Strip characters Excel rejects.
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)
    return text[:SNIPPET_LEN], attachments


def find_sent_folder(imap):
    status, data = imap.list()
    if status != "OK":
        return None
    names = []
    for raw in data:
        line = raw.decode(errors="replace") if isinstance(raw, bytes) else str(raw)
        m = re.match(r'\((?P<flags>[^)]*)\) "?(?P<delim>[^"]*)"? (?P<name>.+)', line)
        if not m:
            continue
        name = m.group("name").strip().strip('"')
        names.append(name)
        if "\\sent" in m.group("flags").lower():
            return name
    lower = {n.lower(): n for n in names}
    for candidate in SENT_FALLBACKS:
        if candidate.lower() in lower:
            return lower[candidate.lower()]
    for n in names:
        if "sent" in n.lower():
            return n
    return None


def quote_folder(name):
    return '"' + name.replace("\\", "\\\\").replace('"', '\\"') + '"'


def fetch_folder(imap, folder, label, since=None, limit=None, with_body=True):
    status, _ = imap.select(quote_folder(folder), readonly=True)
    if status != "OK":
        print(f"  ! Could not open folder '{folder}', skipping.")
        return []

    criteria = ["ALL"]
    if since:
        criteria = ["SINCE", since.strftime("%d-%b-%Y")]
    status, data = imap.uid("SEARCH", None, *criteria)
    if status != "OK":
        return []
    uids = data[0].split()
    if limit:
        uids = uids[-limit:]  # newest N
    print(f"  {label}: {len(uids)} message(s) in '{folder}'")

    part = "BODY.PEEK[]" if with_body else "BODY.PEEK[HEADER]"
    rows = []
    for i in range(0, len(uids), BATCH_SIZE):
        chunk = b",".join(uids[i:i + BATCH_SIZE])
        status, data = imap.uid("FETCH", chunk, f"(FLAGS RFC822.SIZE {part})")
        if status != "OK":
            continue
        for item in data:
            if not isinstance(item, tuple):
                continue
            meta = item[0].decode(errors="replace")
            msg = email.message_from_bytes(item[1])
            size_m = re.search(r"RFC822\.SIZE (\d+)", meta)
            flags_m = re.search(r"FLAGS \(([^)]*)\)", meta)
            size = int(size_m.group(1)) if size_m else len(item[1])
            seen = "\\Seen" in (flags_m.group(1) if flags_m else "")

            from_name, from_addr = parseaddr(decode_str(msg.get("From")))
            snippet, attachments = ("", [])
            if with_body:
                snippet, attachments = extract_text_and_attachments(msg)

            rows.append({
                "Folder": label,
                "Date": parse_date(msg.get("Date")),
                "From Name": from_name,
                "From Email": from_addr.lower(),
                "To": format_addresses(msg.get("To")),
                "Cc": format_addresses(msg.get("Cc")),
                "Subject": decode_str(msg.get("Subject")) or "(no subject)",
                "Attachments": len(attachments),
                "Attachment Names": "; ".join(attachments),
                "Size (KB)": round(size / 1024, 1),
                "Read": "Yes" if seen else "No",
                "Preview": snippet,
                "Message-ID": decode_str(msg.get("Message-ID")),
            })
        print(f"    fetched {min(i + BATCH_SIZE, len(uids))}/{len(uids)}", end="\r")
    print()
    rows.sort(key=lambda r: r["Date"] or datetime.min, reverse=True)
    return rows


# --------------------------------------------------------------------------- Excel

def clean_cell(value):
    if isinstance(value, str):
        value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)
        return value[:32000]
    return value


def write_email_sheet(wb, title, rows):
    ws = wb.create_sheet(title)
    for col, (name, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 22

    for r, row in enumerate(rows, start=2):
        band = BAND_FILL if r % 2 == 0 else None
        for c, (name, _) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=r, column=c, value=clean_cell(row[name]))
            if isinstance(cell.value, str):
                cell.data_type = "s"  # never let email content (e.g. "=...") become a formula
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=name in ("Subject", "Preview", "To"))
            if band:
                cell.fill = band
            if name == "Date":
                cell.number_format = "yyyy-mm-dd hh:mm"
            elif name == "Size (KB)":
                cell.number_format = "#,##0.0"

    ws.freeze_panes = "C2"
    last_col = get_column_letter(len(COLUMNS))
    ws.auto_filter.ref = f"A1:{last_col}{max(len(rows) + 1, 1)}"
    return ws


def write_summary(wb, account, sections, since):
    ws = wb.active
    ws.title = "Summary"
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 16

    ws["A1"] = "Email Export Summary"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Account: {account}"
    ws["A3"] = f"Exported: {datetime.now():%Y-%m-%d %H:%M}"
    ws["A4"] = f"Period: since {since:%Y-%m-%d}" if since else "Period: all messages"
    for ref in ("A2", "A3", "A4"):
        ws[ref].font = BODY_FONT

    def header(row, labels):
        for i, label in enumerate(labels):
            cell = ws.cell(row=row, column=1 + i, value=label)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL

    row = 6
    header(row, ["Folder", "Messages", "Size (MB)"])
    first = row + 1
    for sheet_name in ("Received", "Sent"):
        row += 1
        n = len(sections.get(sheet_name, []))
        ws.cell(row=row, column=1, value=sheet_name).font = BODY_FONT
        # Live formulas so the summary stays in sync if rows are edited.
        ws.cell(row=row, column=2, value=f"=COUNTA('{sheet_name}'!A2:A{n + 1})" if n else 0).font = BODY_FONT
        size_col = get_column_letter([c for c, _ in COLUMNS].index("Size (KB)") + 1)
        c = ws.cell(row=row, column=3, value=f"=SUM('{sheet_name}'!{size_col}2:{size_col}{n + 1})/1024" if n else 0)
        c.font = BODY_FONT
        c.number_format = "#,##0.00"
    row += 1
    ws.cell(row=row, column=1, value="Total").font = Font(name="Arial", bold=True)
    ws.cell(row=row, column=2, value=f"=SUM(B{first}:B{row - 1})").font = Font(name="Arial", bold=True)
    c = ws.cell(row=row, column=3, value=f"=SUM(C{first}:C{row - 1})")
    c.font = Font(name="Arial", bold=True)
    c.number_format = "#,##0.00"

    all_rows = sections.get("Received", []) + sections.get("Sent", [])
    dates = [r["Date"] for r in all_rows if r["Date"]]
    row += 2
    ws.cell(row=row, column=1, value="Oldest message").font = BODY_FONT
    c = ws.cell(row=row, column=2, value=min(dates) if dates else "-")
    c.number_format = "yyyy-mm-dd"
    row += 1
    ws.cell(row=row, column=1, value="Newest message").font = BODY_FONT
    c = ws.cell(row=row, column=2, value=max(dates) if dates else "-")
    c.number_format = "yyyy-mm-dd"

    # Top senders (received) and top recipients (sent)
    senders = Counter(r["From Email"] for r in sections.get("Received", []) if r["From Email"])
    recipients = Counter()
    for r in sections.get("Sent", []):
        for addr in re.findall(r"[\w.+-]+@[\w-]+\.[\w.-]+", r["To"] + " " + r["Cc"]):
            recipients[addr.lower()] += 1

    for title, counter in (("Top 10 Senders (Received)", senders),
                           ("Top 10 Recipients (Sent)", recipients)):
        row += 2
        header(row, [title, "Emails"])
        for addr, count in counter.most_common(10):
            row += 1
            ws.cell(row=row, column=1, value=addr).font = BODY_FONT
            ws.cell(row=row, column=2, value=count).font = BODY_FONT
        if not counter:
            row += 1
            ws.cell(row=row, column=1, value="(none)").font = BODY_FONT

    # Monthly breakdown
    monthly = {}
    for label in ("Received", "Sent"):
        for r in sections.get(label, []):
            if r["Date"]:
                key = r["Date"].strftime("%Y-%m")
                monthly.setdefault(key, {"Received": 0, "Sent": 0})[label] += 1
    row += 2
    header(row, ["Month", "Received", "Sent"])
    for key in sorted(monthly, reverse=True):
        row += 1
        ws.cell(row=row, column=1, value=key).font = BODY_FONT
        ws.cell(row=row, column=2, value=monthly[key]["Received"]).font = BODY_FONT
        ws.cell(row=row, column=3, value=monthly[key]["Sent"]).font = BODY_FONT


# --------------------------------------------------------------------------- main

def main():
    p = argparse.ArgumentParser(description="Export Hostinger sent/received emails to Excel.")
    p.add_argument("--email", required=True, help="Full Hostinger email address")
    p.add_argument("--host", default=IMAP_HOST, help=f"IMAP host (default {IMAP_HOST})")
    p.add_argument("--port", type=int, default=IMAP_PORT, help=f"IMAP SSL port (default {IMAP_PORT})")
    p.add_argument("--since", help="Only export messages on/after this date (YYYY-MM-DD)")
    p.add_argument("--limit", type=int, help="Max number of newest messages per folder")
    p.add_argument("--sent-folder", help="Name of the Sent folder (auto-detected if omitted)")
    p.add_argument("--headers-only", action="store_true",
                   help="Skip message bodies (faster; no preview or attachment names)")
    p.add_argument("--output", help="Output .xlsx path (default emails_<account>_<date>.xlsx)")
    args = p.parse_args()

    since = datetime.strptime(args.since, "%Y-%m-%d") if args.since else None
    password = os.environ.get("HOSTINGER_EMAIL_PASSWORD") or getpass.getpass(f"Password for {args.email}: ")
    output = args.output or f"emails_{args.email.replace('@', '_at_')}_{datetime.now():%Y%m%d}.xlsx"

    print(f"Connecting to {args.host}:{args.port} ...")
    try:
        imap = imaplib.IMAP4_SSL(args.host, args.port)
        imap.login(args.email, password)
    except imaplib.IMAP4.error as e:
        sys.exit(f"Login failed: {e}")
    except OSError as e:
        sys.exit(f"Could not connect to {args.host}:{args.port}: {e}")

    try:
        sent_folder = args.sent_folder or find_sent_folder(imap)
        if not sent_folder:
            print("  ! Sent folder not found; use --sent-folder to specify it.")
        sections = {
            "Received": fetch_folder(imap, "INBOX", "Received", since, args.limit, not args.headers_only),
            "Sent": fetch_folder(imap, sent_folder, "Sent", since, args.limit, not args.headers_only)
            if sent_folder else [],
        }
    finally:
        try:
            imap.logout()
        except Exception:
            pass

    wb = Workbook()
    write_summary(wb, args.email, sections, since)
    write_email_sheet(wb, "Received", sections["Received"])
    write_email_sheet(wb, "Sent", sections["Sent"])
    combined = sorted(sections["Received"] + sections["Sent"],
                      key=lambda r: r["Date"] or datetime.min, reverse=True)
    write_email_sheet(wb, "All Emails", combined)
    wb.save(output)

    print(f"Done. {len(sections['Received'])} received, {len(sections['Sent'])} sent -> {output}")


if __name__ == "__main__":
    main()
